import openpyxl
from openpyxl.styles import Font
from openpyxl.chart import BarChart, Reference

# Create a workbook and select the active sheet
wb = openpyxl.Workbook()
sheet = wb.active
sheet.title = "Sales Data"

# Prepare data
data = [
    ["Product", "January", "February", "March"],
    ["Apples", 50, 20, 35],
    ["Bananas", 80, 63, 24],
    ["Cherries", 75, 82, 13]
]

# Write data to sheet and apply bold font to header
for row_index, row in enumerate(data, start=1):
    for col_index, cell_value in enumerate(row, start=1):
        cell = sheet.cell(row=row_index, column=col_index, value=cell_value)
        
        if row_index == 1:
            cell.font = Font(bold=True)

# Create a Bar Chart
chart = BarChart()
data_range = Reference(sheet, min_col=2, min_row=1, max_col=4, max_row=4)
categories = Reference(sheet, min_col=1, min_row=2, max_row=4)
chart.add_data(data_range, titles_from_data=True)
chart.set_categories(categories)
chart.title = "Sales Data"
chart.x_axis.title = "Products"
chart.y_axis.title = "Sales (in units)"

# Add the chart to the sheet
sheet.add_chart(chart, "F5")

# Save the workbook
filename = "advanced_example.xlsx"
wb.save(filename)
print(f"Workbook saved as {filename}")

